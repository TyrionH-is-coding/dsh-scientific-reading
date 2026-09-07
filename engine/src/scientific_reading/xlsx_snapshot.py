"""从 SQLite 生成长期管理工作簿，并安全回写固定个人字段。"""

from __future__ import annotations

import json
import os
import sys
import shutil
import subprocess
import sqlite3
import tempfile
import re
import time
import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .library_service import library_path
from .library_schema import migrate_library
from .data_guard import root_operation, _file_lock, _key

from .personal_records import USER_FIELDS, normalize
from .xlsx_workbook import XLSX_COLUMNS, REVIEW_COLUMNS, ASSET_COLUMNS, write_workbook
from .xlsx_sync import import_fields, archive_workbook


class XlsxSnapshotService:
    def __init__(self, data_root: Path) -> None:
        self.data_root = Path(data_root).resolve()
        self.target = self.data_root / "library" / "scientific-reading.xlsx"

    @root_operation
    def refresh(self) -> dict[str, Any]:
        with _file_lock(_key(self.data_root), "xlsx", exclusive=True, deadline=time.monotonic() + 30):
            return self._refresh_locked()

    def _refresh_locked(self):
        migrate_library(self.data_root)
        original = hashlib.sha256(self.target.read_bytes()).hexdigest() if self.target.is_file() else None
        imported = {"updated": 0}
        if self.target.is_file():
            imported = import_fields(self)
            if imported.get("status") == "pending":
                return imported
        rows = self._rows()
        review_rows = self._review_rows()
        asset_rows = self._asset_rows()
        self.target.parent.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        try:
            temporary = self._write_temp(rows, review_rows, asset_rows)
            if self._workbook_in_use():
                temporary.unlink(missing_ok=True)
                return self._pending_import("xlsx_in_use", "工作簿正在使用；请保存并关闭表格软件后重试。")
            current = hashlib.sha256(self.target.read_bytes()).hexdigest() if self.target.is_file() else None
            if current != original:
                temporary.unlink(missing_ok=True)
                return self._pending_import("xlsx_changed_during_export", "生成期间工作簿被修改，已保留较新的原表；请关闭后重试。")
            archive_workbook(self)
            os.replace(temporary, self.target)
        except PermissionError as error:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
            self._set_meta("1", "xlsx_replace_permission_denied")
            return {
                "status": "pending",
                "path": str(self.target),
                "error": {"code": "xlsx_replace_permission_denied", "detail": str(error)},
            }
        except (OSError, sqlite3.Error, ValueError) as error:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
            self._set_meta("1", "xlsx_snapshot_failed")
            return {
                "status": "failed",
                "path": str(self.target),
                "error": {"code": "xlsx_snapshot_failed", "detail": str(error)},
            }
        self._set_meta("0", None)
        from .environment_status import EnvironmentStatusService
        status = EnvironmentStatusService(self.data_root)
        with sqlite3.connect(str(library_path(self.data_root))) as conn:
            conn.executemany("UPDATE items SET xlsx_sync_state='ready',xlsx_error=NULL WHERE paper_id=? "
                             "AND coalesce(personal_updated_at,'')=? AND updated_at=?",
                             [(r["文献 ID"], r["个人记录更新时间"], r["处理更新时间"]) for r in rows])
            conn.execute("INSERT OR REPLACE INTO library_meta VALUES('xlsx_last_export',?)", (json.dumps({**self._receipt, "rows": len(rows)}, ensure_ascii=False),))
        status._write(status.snapshot())
        return {"status": "success", "path": str(self.target), "rows": len(rows),
                "updated": imported["updated"], "conflicts": 0, **self._receipt}

    def _review_rows(self):
        from .xlsx_reading import reading_records
        return reading_records(self)

    def _rows(self):
        migrate_library(self.data_root)
        with sqlite3.connect(str(library_path(self.data_root))) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT i.*, f.name AS folder_name, a.rel_path AS pdf_path, a.sha256 AS pdf_sha, "
                "(SELECT rel_path FROM artifacts WHERE paper_id=i.paper_id AND kind='reader' "
                "AND status='ready' LIMIT 1) AS reader_path, "
                "(SELECT GROUP_CONCAT(tag, ', ') FROM item_tags WHERE paper_id=i.paper_id) AS tags "
                "FROM items i LEFT JOIN folders f ON f.folder_id=i.folder_id "
                "LEFT JOIN attachments a ON a.paper_id=i.paper_id ORDER BY i.created_at,i.paper_id"
            ).fetchall()
        result = []
        for row in rows:
            paper_root = self.data_root / "papers" / row["paper_id"]
            reader = ""
            if row["pdf_sha"] and row["reader_path"]:
                reader = self._active_reader_rel(paper_root, paper_root / "generations" / row["pdf_sha"][:16], row["reader_path"])
            if row["reader_path"] == "reading/reader.html":
                from .library_service import LibraryService
                library = LibraryService(self.data_root)
                try:
                    library.validate_reader(row["paper_id"], row["reader_path"])
                    reader = row["reader_path"]
                except (OSError, ValueError):
                    pass
                finally:
                    library.close()
            processing = "Reader 可用" if reader else "缺 PDF" if not row["pdf_path"] else "待精读"
            if row["last_error"]:
                processing = "需处理"
            paper = {"文献名": row["title"], "分类": row["folder_name"] or "未分类", "PDF": row["pdf_path"] or "",
                     "Reader": reader, "年份": row["year"], "期刊": row["journal"], "标签": row["tags"] or "",
                     "处理进度": processing, "个人记录更新时间": row["personal_updated_at"] or "",
                     "作者": "; ".join(json.loads(row["authors_json"] or "[]")), "DOI": row["doi"], "PMID": row["pmid"],
                     "Abstract (EN)": row["abstract_en"], "Abstract (ZH)": row["abstract_zh"], "文献链接": row["source_url"],
                     "入库时间": row["created_at"], "处理更新时间": row["updated_at"], "文献 ID": row["paper_id"]}
            paper.update({label: normalize(field, row[field]) for label, field in USER_FIELDS.items()})
            result.append(paper)
        return result

    def _asset_rows(self) -> list[tuple[Any, ...]]:
        migrate_library(self.data_root)
        conn = sqlite3.connect(str(library_path(self.data_root)))
        conn.row_factory = sqlite3.Row
        try:
            papers = list(
                conn.execute(
                    "SELECT i.paper_id, i.title, a.rel_path AS pdf_path, "
                    "a.sha256 AS source_sha256, "
                    "(SELECT rel_path FROM artifacts WHERE paper_id=i.paper_id "
                    "AND kind IN ('reader','full_read_html','full_read') "
                    "AND status='ready' ORDER BY updated_at DESC LIMIT 1) AS html_path "
                    "FROM items i LEFT JOIN attachments a ON a.paper_id=i.paper_id "
                    "ORDER BY i.created_at, i.paper_id"
                )
            )
        finally:
            conn.close()
        result: list[tuple[Any, ...]] = []
        for paper in papers:
            source_sha = paper["source_sha256"]
            if not isinstance(source_sha, str) or not re.fullmatch(r"[0-9a-f]{64}", source_sha):
                continue
            paper_root = self.data_root / "papers" / paper["paper_id"]
            generation_root = paper_root / "generations" / source_sha[:16]
            manifest = self._load_json(generation_root / "manifest.json")
            assets = manifest.get("assets") if isinstance(manifest, dict) else None
            if not isinstance(assets, list):
                continue
            generation_prefix = generation_root.relative_to(paper_root).as_posix()
            reader_rel = self._active_reader_rel(
                paper_root, generation_root, paper["html_path"]
            )
            reader_manifest = self._load_json(
                generation_root / "reading" / "reader-manifest.json"
            )
            caption_blocks: dict[str, str] = {}
            for asset in reader_manifest.get("assets", []) if isinstance(reader_manifest, dict) else []:
                if not isinstance(asset, dict):
                    continue
                logical_id = self._logical_asset_id(asset.get("id"), asset.get("path"))
                block_id = asset.get("caption_block_id")
                if logical_id and isinstance(block_id, str) and block_id:
                    caption_blocks.setdefault(logical_id, block_id)
            translations = self._load_json(
                generation_root / "reading" / "full" / "translations.json"
            )
            translated_captions = {
                item["block_id"]: item["translation_zh"]
                for item in translations.get("translations", [])
                if isinstance(item, dict)
                and isinstance(item.get("block_id"), str)
                and isinstance(item.get("translation_zh"), str)
            } if isinstance(translations, dict) else {}
            grouped: dict[str, dict[str, Any]] = {}
            for asset in assets:
                if not isinstance(asset, dict):
                    continue
                rel_path = self._safe_generation_file(
                    generation_root, asset.get("relative_path")
                )
                logical_id = self._logical_asset_id(
                    asset.get("asset_id"), asset.get("relative_path")
                )
                kind = asset.get("kind")
                if not logical_id or kind not in {"figure", "table"} or not rel_path:
                    continue
                row = grouped.setdefault(
                    logical_id,
                    {
                        "kind": kind,
                        "page": asset.get("page"),
                        "caption": asset.get("caption") or "",
                        "image": "",
                        "html": "",
                    },
                )
                suffix = Path(rel_path).suffix.lower()
                paper_rel = f"{generation_prefix}/{rel_path}"
                if suffix == ".html":
                    row["html"] = paper_rel
                else:
                    row["image"] = paper_rel
                structured = self._safe_generation_file(
                    generation_root, asset.get("structured_path")
                )
                if structured:
                    row["html"] = f"{generation_prefix}/{structured}"
                if not row["caption"] and asset.get("caption"):
                    row["caption"] = asset["caption"]
                if not row["page"] and asset.get("page"):
                    row["page"] = asset["page"]
            for asset_id, asset in grouped.items():
                caption_block = caption_blocks.get(asset_id, "")
                reader_locator = (
                    f"{reader_rel}#block-{caption_block}"
                    if reader_rel and caption_block else ""
                )
                result.append((
                    paper["paper_id"], paper["title"], asset_id,
                    "图" if asset["kind"] == "figure" else "表",
                    asset["page"], translated_captions.get(caption_block, ""),
                    asset["caption"], asset["image"], asset["html"], reader_locator,
                ))
        return result

    @staticmethod
    def _load_json(path: Path) -> dict[str, Any]:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return {}
        return payload if isinstance(payload, dict) else {}

    @staticmethod
    def _logical_asset_id(asset_id: Any, rel_path: Any) -> str:
        if not isinstance(asset_id, str) or not asset_id:
            return ""
        if isinstance(rel_path, str) and Path(rel_path).suffix.lower() == ".html":
            return asset_id.removesuffix("-html")
        return asset_id

    @staticmethod
    def _safe_generation_file(generation_root: Path, rel_path: Any) -> str:
        if not isinstance(rel_path, str) or not rel_path or Path(rel_path).is_absolute():
            return ""
        root = generation_root.resolve()
        candidate = (root / rel_path).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            return ""
        return Path(rel_path).as_posix() if candidate.is_file() else ""

    @staticmethod
    def _active_reader_rel(
        paper_root: Path, generation_root: Path, rel_path: Any
    ) -> str:
        if not isinstance(rel_path, str) or not rel_path or Path(rel_path).is_absolute():
            return ""
        root = paper_root.resolve()
        generation = generation_root.resolve()
        candidate = (root / rel_path).resolve()
        try:
            candidate.relative_to(generation)
        except ValueError:
            return ""
        return Path(rel_path).as_posix() if candidate.is_file() else ""

    def _write_temp(self, rows, review_rows, asset_rows) -> Path:
        handle, name = tempfile.mkstemp(prefix=".scientific-reading-", suffix=".xlsx", dir=self.target.parent)
        os.close(handle)
        path = Path(name)
        try:
            self._receipt = write_workbook(self, path, rows, review_rows, asset_rows)
            with path.open("r+b") as stream:
                os.fsync(stream.fileno())
        except Exception:
            path.unlink(missing_ok=True)
            raise
        return path

    def _paper_link(
        self, paper_id: Any, rel_path: Any, fragment: str | None = None
    ) -> str:
        if not isinstance(paper_id, str) or not isinstance(rel_path, str) or not rel_path:
            return ""
        paper_root = (self.data_root / "papers" / paper_id).resolve()
        if Path(rel_path).is_absolute():
            return ""
        candidate = (paper_root / rel_path).resolve()
        try:
            candidate.relative_to(paper_root)
        except ValueError:
            return ""
        if not candidate.exists():
            return ""
        target = os.path.relpath(candidate, self.target.parent).replace(os.sep, "/")
        return f"{target}#{fragment}" if fragment else target

    @root_operation
    def import_user_fields(self):
        migrate_library(self.data_root)
        with _file_lock(_key(self.data_root), "xlsx", exclusive=True, deadline=time.monotonic() + 30):
            return import_fields(self)

    def _workbook_in_use(self) -> bool:
        # Unix permits replacing open files; Office owner files also protect that path.
        return any(self.target.with_name(name).exists() for name in (
            "~$" + self.target.name, ".~lock." + self.target.name + "#",
        ))

    def _pending_import(self, code: str, detail: str) -> dict[str, Any]:
        self._set_meta("1", code)
        return {
            "status": "pending",
            "path": str(self.target),
            "error": {"code": code, "detail": detail},
        }

    def locate(self, paper_id: str, *, opener=None) -> dict[str, Any]:
        if not isinstance(paper_id, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", paper_id) or ".." in paper_id:
            raise ValueError("paper_id_invalid")
        if not self.target.is_file():
            raise ValueError("xlsx_not_generated")
        workbook = load_workbook(self.target, read_only=True)
        try:
            sheet = workbook["文献"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            column = headers.get("文献 ID")
            row_number = next((row for row in range(2, sheet.max_row + 1) if sheet.cell(row, column).value == paper_id), None) if column else None
        finally:
            workbook.close()
        if row_number is None:
            raise ValueError("xlsx_identity_not_found")
        selected = (opener or self._open_excel)(self.target, row_number)
        return {"status": "opened", "paper_id": paper_id, "row": row_number, "selected": bool(selected)}

    @staticmethod
    def _open_excel(path: Path, row_number: int) -> bool:
        if sys.platform != "win32":
            command = "/usr/bin/open" if sys.platform == "darwin" else shutil.which("xdg-open")
            if not command:
                raise ValueError("spreadsheet_opener_unavailable")
            try:
                subprocess.run([command, str(path)], check=True, timeout=15,
                               stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                               stderr=subprocess.DEVNULL)
            except (OSError, subprocess.SubprocessError) as error:
                raise ValueError("spreadsheet_open_failed") from error
            return False
        try:
            import win32com.client  # type: ignore[import-not-found]

            excel = win32com.client.Dispatch("Excel.Application")
            workbook = excel.Workbooks.Open(str(path))
            workbook.Worksheets("文献").Cells(row_number, 1).Select()
            excel.Visible = True
            return True
        except (ImportError, OSError):
            if os.name != "nt":
                raise ValueError("excel_unavailable")
            os.startfile(path)  # type: ignore[attr-defined]
            return False

    def _set_meta(self, pending: str, error: str | None) -> None:
        conn = sqlite3.connect(str(library_path(self.data_root)))
        try:
            conn.execute("INSERT OR REPLACE INTO library_meta(key,value) VALUES('xlsx_pending',?)", (pending,))
            if error:
                conn.execute("INSERT OR REPLACE INTO library_meta(key,value) VALUES('xlsx_error',?)", (error,))
            else:
                conn.execute("DELETE FROM library_meta WHERE key='xlsx_error'")
            conn.commit()
        finally:
            conn.close()
