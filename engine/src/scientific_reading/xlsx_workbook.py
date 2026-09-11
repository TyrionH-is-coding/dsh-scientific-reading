"""固定的长期管理工作簿；系统字段仅展示，个人字段在同步时校验。"""

from collections import Counter
from datetime import UTC, datetime
import json
from pathlib import Path
import re
import sqlite3
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .personal_records import USER_FIELDS, READING_STATES
from .library_views import xlsx_columns

XLSX_COLUMNS = (
    "文献名", "分类", "阅读进度", "与课题的关系", "下一步", "PDF", "Reader", "阅读成果", "图表索引",
    "年份", "期刊", "标签", "处理进度", "个人记录更新时间", "个人理解程度", "个人思考", "用户笔记",
    "作者", "DOI", "PMID", "Abstract (EN)", "Abstract (ZH)", "文献链接", "入库时间", "处理更新时间",
    "文献 ID", "行标识",
)
REVIEW_COLUMNS = ("文献名", "记录类别", "内容", "依据类型", "确认状态", "证据位置", "打开原文", "更新时间", "文献 ID", "记录 ID")
ASSET_COLUMNS = ("文献名", "原文图表编号", "类型", "PDF 页码", "图注", "关联阅读结论", "原图", "结构表格", "打开原文", "状态", "源文图注", "文献 ID", "资产 ID")


def _text(cell, value):
    cell.value = value
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError("xlsx_cell_text_too_long")
        cell.data_type = "s"  # 论文和笔记里的等号不是可执行公式。


def _sheet(workbook, name, columns, rows, table_name, widths):
    sheet = workbook.create_sheet(name)
    sheet.append(columns)
    for number, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            _text(sheet.cell(number, column), value)
        sheet.row_dimensions[number].height = 68 if name == "文献" else 92
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="183D42")
        cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, header in enumerate(columns, 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths.get(header, 16)
        for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for cell in cells:
                cell.font = Font(name="Calibri", color="233F43", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header in USER_FIELDS:
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")
        if header in {"文献 ID", "行标识", "记录 ID", "资产 ID", "字段 ID"}:
            sheet.column_dimensions[letter].hidden = True
    reference = f"A1:{get_column_letter(len(columns))}{max(1, sheet.max_row)}"
    if rows:
        # 同一区域叠加 worksheet autoFilter 会使原生 Excel 拒绝打开。
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    # Excel 不允许对锁定单元格的保护表正常排序；回写白名单负责数据保护。
    sheet.protection.sheet = False
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def _jump(cell, paper_id, sheet, columns, label):
    col = get_column_letter(columns.index("文献 ID") + 1)
    key = paper_id.replace('"', '""')
    # MATCH 按稳定键定位，目标表经用户排序后入口仍指向原论文。
    cell.value = f'=HYPERLINK("#\'{sheet}\'!A"&MATCH("{key}",\'{sheet}\'!${col}:${col},0),"{label}")'
    cell.font = Font(color="19736C", underline="single")


def _link(service, cell, paper_id, relative, label=None):
    path, _, fragment = (relative or "").partition("#")
    target = service._paper_link(paper_id, path, fragment or None)
    if target:
        cell.hyperlink = target
        if label:
            _text(cell, label)
        cell.font = Font(color="19736C", underline="single")
    else:
        _text(cell, "未就绪")


def write_workbook(service, path, papers, records, assets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    export_id = uuid.uuid4().hex
    stamp = datetime.now(UTC).isoformat()
    with sqlite3.connect(service.data_root / "library.sqlite") as conn:
        columns = xlsx_columns(conn, getattr(service, "folder_id", None) or "")
        research = {(row[0], row[1]): {"value": row[2], "origin": row[3], "evidence": json.loads(row[4]), "revision": row[5], "updated_at": row[6]}
                    for row in conn.execute("SELECT * FROM research_values")}
    keys = [column["key"] for column in columns]
    labels = [column["label"] for column in columns]
    if len(set(labels)) != len(labels):
        raise ValueError("display_column_label_duplicate")
    baseline = {}
    rows = []
    for paper in papers:
        key = uuid.uuid4().hex
        paper["行标识"] = key
        custom = {column["field_id"]: research.get((paper["文献 ID"], column["field_id"]), {"value": "", "revision": 0, "origin": "manual", "evidence": []}) for column in columns if column["custom"]}
        baseline[key] = {"paper_id": paper["文献 ID"], "fields": {field: paper[label] for label, field in USER_FIELDS.items()}, "custom": custom}
        paper.update({field: value["value"] for field, value in custom.items()})
        rows.append(tuple(paper.get(header, "") for header in keys))
    widths = {"文献名": 44, "与课题的关系": 34, "下一步": 30, "期刊": 24, "个人记录更新时间": 25,
              "个人思考": 44, "用户笔记": 48, "个人理解程度": 24, "Abstract (EN)": 65, "Abstract (ZH)": 65,
              "作者": 30, "DOI": 32, "文献链接": 40, "内容": 70, "证据位置": 30, "更新时间": 25,
              "图注": 64, "源文图注": 64, "关联阅读结论": 64, "状态": 26}
    sheet = _sheet(workbook, "文献", labels, rows, "Literature", {column["label"]: column.get("width", widths.get(column["key"], 30 if column["custom"] else 16)) for column in columns})
    for index, column in enumerate(columns, 1):
        dim = sheet.column_dimensions[get_column_letter(index)]
        dim.hidden = column.get("hidden", column["field_id"] in {"paper_id", "row_token", "authors", "doi", "pmid", "abstract_en", "abstract_zh", "source_url", "created_at", "updated_at"})
        for cells in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=column.get("wrap", True))
                if column["editable"]:
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")
    review_rows = [tuple(record.get(header, "") for header in REVIEW_COLUMNS) for record in records]
    reading = _sheet(workbook, "阅读成果", REVIEW_COLUMNS, review_rows, "ReadingRecords", widths)
    asset_rows = []
    for asset in assets:
        paper_id, title, asset_id, kind, page, zh, caption, image, html, reader = asset
        match = re.search(r"(?:Figure|Fig\.?|Table|图|表)\s*\d+[A-Za-z]?", caption or "", re.IGNORECASE)
        block = reader.partition("#block-")[2]
        linked = [r["内容"] for r in records if r["文献 ID"] == paper_id and block
                  and block in r.get("block_ids", []) and r.get("current")]
        asset_rows.append((title, match.group() if match else "图号待核对", kind, page, zh or caption or "图注未提供",
                           "\n".join(linked) or "未建立解读关联", image, html, reader,
                           "原图可用" if image else "结构表可用", caption, paper_id, asset_id))
    figures = _sheet(workbook, "图表索引", ASSET_COLUMNS, asset_rows, "FigureIndex", widths)
    paper_by_id = {p["文献 ID"]: p for p in papers}
    count_records = Counter(r["文献 ID"] for r in records)
    count_assets = Counter(a[0] for a in assets)
    for row, paper in enumerate(papers, 2):
        paper_id = paper["文献 ID"]
        for name in ("PDF", "Reader"):
            _link(service, sheet.cell(row, keys.index(name) + 1), paper_id, paper[name], "打开")
        for name, target_columns, counts in (("阅读成果", REVIEW_COLUMNS, count_records), ("图表索引", ASSET_COLUMNS, count_assets)):
            cell = sheet.cell(row, keys.index(name) + 1)
            if counts[paper_id]:
                _jump(cell, paper_id, name, target_columns, f"查看 {counts[paper_id]} 项")
            else:
                _text(cell, "未生成")
        cell = sheet.cell(row, keys.index("文献链接") + 1)
        if isinstance(cell.value, str) and cell.value.startswith(("https://", "http://")):
            cell.hyperlink = cell.value
    for row, record in enumerate(records, 2):
        paper_id = record["文献 ID"]
        _jump(reading.cell(row, 1), paper_id, "文献", keys, record["文献名"].replace('"', '""'))
        _link(service, reading.cell(row, 7), paper_id, record.get("打开原文"), "回到证据")
    for row, asset in enumerate(assets, 2):
        paper_id, title, _, _, page, _, _, image, html, reader = asset
        _jump(figures.cell(row, 1), paper_id, "文献", keys, title.replace('"', '""'))
        for column, relative, label in ((7, image, "打开原图"), (8, html, "打开表格"), (9, reader, "回到原文")):
            _link(service, figures.cell(row, column), paper_id, relative, label)
        pdf = paper_by_id[paper_id]["PDF"]
        if page and pdf:
            _link(service, figures.cell(row, 4), paper_id, f"{pdf}#page={page}")
    validation = DataValidation(type="list", formula1='"' + ','.join(READING_STATES) + '"', allow_blank=False)
    validation.errorTitle = "请选择阅读进度"
    validation.error = "请选择未读、在读、已读或待复读。"
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    sheet.add_data_validation(validation)
    state_column = get_column_letter(keys.index("阅读进度") + 1)
    validation.add(f"{state_column}2:{state_column}{max(2, sheet.max_row)}")
    if research:
        research_rows = []
        for paper in papers:
            for column in columns:
                value = research.get((paper["文献 ID"], column["field_id"]))
                if value is None:
                    continue
                evidence = "\n".join((f"摘要 {entry['language']}：{entry['quote']}" if entry.get("kind") == "abstract"
                                      else f"PDF 第 {entry['page']} 页 · {entry['block_id']}：{entry['quote']}") for entry in value["evidence"])
                research_rows.append((paper["文献名"], column["label"], value["value"], "手动" if value["origin"] == "manual" else "AI（待核对）", evidence,
                                      value["updated_at"], paper["文献 ID"], column["field_id"]))
        _sheet(workbook, "研究字段", ("文献名", "字段", "内容", "来源", "证据", "更新时间", "文献 ID", "字段 ID"), research_rows, "ResearchFields", {**widths, "证据": 65})
    instructions = workbook.create_sheet("说明")
    for line in (
        "Deep Literature · 文献管理", "黄色个人字段与自定义研究字段可编辑；其余信息由系统维护，修改后不会回写。",
        "在文献工作台保存列名、顺序、隐藏、宽度和换行方案；Excel 内可整列重排，直接改表头时同步会暂停并保留原表。",
        "研究字段表展示手动或 AI 来源和证据；编辑文献表内的研究字段会改为手动内容。AI 不会自动覆盖手动内容。",
        "保存并关闭 Excel / LibreOffice / Numbers 后，在 Codex 对话中说“同步文献表”。",
        "使用整张表的筛选或排序。不要删除文献行、身份列或只移动部分单元格。",
        "阅读进度由你维护；生成 Reader 不会自动标记为已读。",
        "与课题的关系：为何收藏。下一步：后续行动。个人理解程度：尚不理解之处。",
        "个人思考保留判断与假设；用户笔记保留摘录、备忘。旧字段不会合并或删除。",
        "阅读成果区分 AI 待核对导读与用户确认记录；证据定位有效不等于结论已被证明。",
        "图表索引按当前 PDF 生成。同一表的原图与 HTML 共用一行；没有证据关联时不自动解释图意。",
        "同步冲突时保留原表和数据库内容；核对具体论文与字段后，再明确选择保留内容。",
        "重写前原工作簿按内容哈希保存在 backups/xlsx；备份整个文献根目录可保留库、基线、原文与成果。",
        f"本次导出：{stamp}（UTC）；文献 {len(papers)} 篇，阅读记录 {len(records)} 条，图表 {len(assets)} 项。",
        "Windows Excel 可定位行；macOS / Linux 用默认表格软件打开，并在返回结果中显示行号。",
    ):
        instructions.append((line,))
    instructions.column_dimensions["A"].width = 110
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="center")
        instructions.row_dimensions[row[0].row].height = 36
    sync = workbook.create_sheet("_同步")
    sync.append(("contract", "xlsx-library-v3"))
    sync.append(("export_id", export_id))
    sync.append(("columns_json", json.dumps(columns, ensure_ascii=False)))
    sync.append(("scope_folder_id", getattr(service, "folder_id", None) or ""))
    sync.sheet_state = "veryHidden"
    try:
        workbook.save(path)
    finally:
        workbook.close()
    with sqlite3.connect(service.data_root / "library.sqlite") as conn:
        conn.execute("INSERT INTO xlsx_exports VALUES(?,?,?)", (export_id, stamp, json.dumps({"rows": baseline, "columns": columns, "scope_folder_id": getattr(service, "folder_id", None) or ""}, ensure_ascii=False)))
    return {"export_id": export_id, "exported_at": stamp}
