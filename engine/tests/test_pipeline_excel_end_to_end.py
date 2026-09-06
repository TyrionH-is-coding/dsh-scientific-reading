from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import time
from pathlib import Path

import openpyxl

from scientific_reading.background_launcher import BackgroundLauncher
from scientific_reading.background_store import BackgroundJobStore
from scientific_reading.library_service import LibraryService, library_path
from scientific_reading.models import PaperMetadata, StageRecord
from scientific_reading.reading_pipeline import ReadingPipeline
from scientific_reading.reading_pipeline_models import ReadingPipelineState
from scientific_reading.workspace import PaperWorkspace
from scientific_reading.xlsx_snapshot import XlsxSnapshotService


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _publish_local_reader(data_root: Path) -> tuple[str, Path]:
    metadata = PaperMetadata(
        title="虚构材料响应研究",
        authors=["测试作者"],
        doi="10.5555/fictional.xlsx.pipeline",
        year=2026,
        journal="离线验证期刊",
        source_url="https://publisher.invalid/fictional-paper",
    )
    library = LibraryService(data_root)
    try:
        paper_id = library.ingest(metadata)["paper_id"]
    finally:
        library.close()

    workspace = PaperWorkspace.create_for_paper_id(data_root, paper_id, metadata)
    workspace.source_pdf.write_bytes(b"%PDF-1.4\n% offline fixture\n%%EOF\n")
    source_sha = _sha256(workspace.source_pdf)
    parser = workspace.parsed_dir / "mineru" / "source_map.json"
    parser.parent.mkdir(parents=True, exist_ok=True)
    parser.write_text(
        json.dumps({"contract": "offline-mineru-fixture-v1"}), encoding="utf-8"
    )
    translation = workspace.reading_dir / "full" / "translations.json"
    translation.parent.mkdir(parents=True, exist_ok=True)
    translation.write_text(
        json.dumps({"contract": "offline-translation-fixture-v1"}), encoding="utf-8"
    )
    workspace.reader_html.write_text(
        "<!doctype html><html><body>离线精读夹具</body></html>", encoding="utf-8"
    )
    workspace.reader_manifest.write_text(
        json.dumps(
            {
                "contract": "reader-manifest-v1",
                "paper_id": paper_id,
                "source_pdf_sha256": source_sha,
                "parser_manifest_sha256": _sha256(parser),
                "translation_manifest_sha256": _sha256(translation),
                "reader_sha256": _sha256(workspace.reader_html),
                "generated_at": "2026-09-05T00:00:00+00:00",
                "source_blocks": [
                    {
                        "block_id": "block-1",
                        "page": 1,
                        "source_type": "text",
                        "source_index": 0,
                    }
                ],
                "assets": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    job = workspace.load_job()
    job.stages["paper_parse_upgrade"] = StageRecord(
        status="completed", result={"source_sha256": source_sha}
    )
    workspace.save_job(job)

    library = LibraryService(data_root)
    try:
        library.record_pdf_attachment(
            paper_id, source_sha, workspace.source_pdf.stat().st_size
        )
        library.publish_reader(paper_id, "reading/reader.html")
    finally:
        library.close()
    return paper_id, workspace.reader_html


def _schedule_real_xlsx_worker(
    data_root: Path, paper_id: str, parent_job_id: str
) -> dict:
    state = ReadingPipelineState(
        paper_id=paper_id,
        parent_job_id=parent_job_id,
        current_stage="schedule_derived_updates",
        source_pdf_sha256="a" * 64,
        reader_source_sha256="a" * 64,
    )
    launched = ReadingPipeline(data_root)._default_stage_runner(
        "schedule_derived_updates", state, None
    )
    return _wait_for_terminal_status(data_root, launched["xlsx_job_id"])


def _wait_for_terminal_status(data_root: Path, job_id: str) -> dict:
    store = BackgroundJobStore(data_root)
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        status = store.load_status(job_id)
        if status.state in {
            "completed",
            "waiting_agent",
            "waiting_user",
            "failed",
            "interrupted",
        }:
            return status.to_dict()
        time.sleep(0.05)
    raise AssertionError(f"xlsx worker did not finish: {status.to_dict()}")


def _wait_for_worker_exit(data_root: Path, job_id: str) -> None:
    store = BackgroundJobStore(data_root)
    launch_path = store.handle(job_id).root / "launch.json"
    marker = json.loads(launch_path.read_text(encoding="utf-8"))
    pid = marker["pid"]
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline and store._pid_is_alive(pid):
        time.sleep(0.05)
    assert not store._pid_is_alive(pid), f"xlsx worker process still alive: {pid}"


def _row_for_paper(workbook, paper_id: str) -> tuple[object, int, dict[str, int]]:
    sheet = workbook["文献"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    row_number = next(
        row
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, headers["文献 ID"]).value == paper_id
    )
    return sheet, row_number, headers


def test_real_scheduled_worker_round_trips_only_user_owned_fields(
    tmp_path, monkeypatch
) -> None:
    engine_src = Path(__file__).resolve().parents[1] / "src"
    inherited = os.environ.get("PYTHONPATH")
    monkeypatch.setenv(
        "PYTHONPATH",
        str(engine_src) + (os.pathsep + inherited if inherited else ""),
    )
    paper_id, _reader = _publish_local_reader(tmp_path)

    first = _schedule_real_xlsx_worker(
        tmp_path, paper_id, "job_excel_parent_0001"
    )

    assert first["state"] == "completed", first
    assert first["result"]["status"] == "success"
    snapshot = XlsxSnapshotService(tmp_path)
    workbook = openpyxl.load_workbook(snapshot.target)
    sheet, row_number, headers = _row_for_paper(workbook, paper_id)
    sheet.cell(row_number, headers["个人思考"], "保留个人思考")
    sheet.cell(row_number, headers["个人理解程度"], "深入")
    sheet.cell(row_number, headers["用户笔记"], "保留用户笔记")
    sheet.cell(row_number, headers["文献名"], "不得回写的伪造标题")
    sheet.cell(row_number, headers["阅读状态"], "不得回写的伪造状态")
    workbook.save(snapshot.target)
    workbook.close()

    second = _schedule_real_xlsx_worker(
        tmp_path, paper_id, "job_excel_parent_0002"
    )

    assert second["state"] == "completed", second
    assert second["result"]["status"] == "success"
    with sqlite3.connect(library_path(tmp_path)) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT title, status, personal_thoughts, understanding_level, user_notes "
            "FROM items WHERE paper_id=?",
            (paper_id,),
        ).fetchone()
        meta = dict(conn.execute("SELECT key, value FROM library_meta"))
    assert dict(row) == {
        "title": "虚构材料响应研究",
        "status": "full_read_ready",
        "personal_thoughts": "保留个人思考",
        "understanding_level": "深入",
        "user_notes": "保留用户笔记",
    }
    assert meta["xlsx_pending"] == "0"
    assert "xlsx_error" not in meta

    workbook = openpyxl.load_workbook(snapshot.target, read_only=True)
    sheet, row_number, headers = _row_for_paper(workbook, paper_id)
    assert sheet.cell(row_number, headers["文献名"]).value == "虚构材料响应研究"
    assert sheet.cell(row_number, headers["阅读状态"]).value == "full_read_ready"
    assert sheet.cell(row_number, headers["个人思考"]).value == "保留个人思考"
    assert sheet.cell(row_number, headers["个人理解程度"]).value == "深入"
    assert sheet.cell(row_number, headers["用户笔记"]).value == "保留用户笔记"
    workbook.close()


def test_published_reader_is_present_as_a_valid_excel_path(tmp_path, monkeypatch) -> None:
    engine_src = Path(__file__).resolve().parents[1] / "src"
    inherited = os.environ.get("PYTHONPATH")
    monkeypatch.setenv(
        "PYTHONPATH",
        str(engine_src) + (os.pathsep + inherited if inherited else ""),
    )
    paper_id, reader = _publish_local_reader(tmp_path)
    with sqlite3.connect(library_path(tmp_path)) as conn:
        conn.execute(
            "INSERT INTO artifacts(paper_id, kind, rel_path, status, updated_at) "
            "VALUES(?, 'full_read_html', 'reading/obsolete.html', 'stale', ?) ",
            (paper_id, "2099-01-01T00:00:00+00:00"),
        )

    status = _schedule_real_xlsx_worker(
        tmp_path, paper_id, "job_excel_parent_html"
    )

    assert status["state"] == "completed", status
    workbook = openpyxl.load_workbook(
        XlsxSnapshotService(tmp_path).target, read_only=True
    )
    sheet, row_number, headers = _row_for_paper(workbook, paper_id)
    stored_path = sheet.cell(row_number, headers["精读 HTML"]).value
    workbook.close()
    assert stored_path
    candidate = Path(stored_path)
    if not candidate.is_absolute():
        candidate = tmp_path / "papers" / paper_id / candidate
    assert candidate.resolve() == reader.resolve()
    assert candidate.is_file()


def test_real_scheduled_worker_preserves_notes_across_pending_and_retry(
    tmp_path, monkeypatch
) -> None:
    engine_src = Path(__file__).resolve().parents[1] / "src"
    inherited = os.environ.get("PYTHONPATH")
    monkeypatch.setenv(
        "PYTHONPATH",
        str(engine_src) + (os.pathsep + inherited if inherited else ""),
    )
    paper_id, _reader = _publish_local_reader(tmp_path)
    ready = _schedule_real_xlsx_worker(
        tmp_path, paper_id, "job_excel_parent_pending_seed"
    )
    assert ready["state"] == "completed", ready

    snapshot = XlsxSnapshotService(tmp_path)
    workbook = openpyxl.load_workbook(snapshot.target)
    sheet, row_number, headers = _row_for_paper(workbook, paper_id)
    sheet.cell(row_number, headers["文献 ID"], "changed_identity")
    sheet.cell(row_number, headers["用户笔记"], "身份修复前必须保留")
    workbook.save(snapshot.target)
    workbook.close()
    pending_bytes = snapshot.target.read_bytes()

    pending = _schedule_real_xlsx_worker(
        tmp_path, paper_id, "job_excel_parent_pending_retry"
    )

    assert pending["state"] == "waiting_user", pending
    assert pending["reason_code"] == "xlsx_identity_conflict"
    assert snapshot.target.read_bytes() == pending_bytes
    with sqlite3.connect(library_path(tmp_path)) as conn:
        meta = dict(conn.execute("SELECT key, value FROM library_meta"))
    assert meta["xlsx_pending"] == "1"
    assert meta["xlsx_error"] == "xlsx_identity_conflict"

    workbook = openpyxl.load_workbook(snapshot.target)
    sheet, row_number, headers = _row_for_paper(workbook, "changed_identity")
    sheet.cell(row_number, headers["文献 ID"], paper_id)
    workbook.save(snapshot.target)
    workbook.close()
    store = BackgroundJobStore(tmp_path)
    _wait_for_worker_exit(tmp_path, pending["job_id"])
    store.transition(pending["job_id"], "queued")
    BackgroundLauncher(tmp_path).launch_existing(pending["job_id"])
    recovered = _wait_for_terminal_status(tmp_path, pending["job_id"])

    assert recovered["state"] == "completed", recovered
    with sqlite3.connect(library_path(tmp_path)) as conn:
        note = conn.execute(
            "SELECT user_notes FROM items WHERE paper_id=?", (paper_id,)
        ).fetchone()[0]
        meta = dict(conn.execute("SELECT key, value FROM library_meta"))
    assert note == "身份修复前必须保留"
    assert meta["xlsx_pending"] == "0"
    assert "xlsx_error" not in meta
