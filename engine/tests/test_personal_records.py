import sqlite3

import openpyxl
import pytest

from scientific_reading.library_service import LibraryService
from scientific_reading.models import PaperMetadata
from scientific_reading.xlsx_snapshot import XlsxSnapshotService


def seed(root):
    library = LibraryService(root)
    ids = [library.ingest(PaperMetadata(title=f"测试论文 {i}"))["paper_id"] for i in range(2)]
    library.close()
    return ids


def edit(service, values, *, reverse=False):
    book = openpyxl.load_workbook(service.target)
    sheet = book["文献"]
    headers = {c.value: c.column for c in sheet[1]}
    for name, value in values.items():
        sheet.cell(2, headers[name]).value = value
    if reverse:
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        for number, row in enumerate(reversed(rows), 2):
            for column, value in enumerate(row, 1):
                sheet.cell(number, column).value = value
    book.save(service.target)
    book.close()


def record(root, paper_id):
    with sqlite3.connect(root / "library.sqlite") as conn:
        conn.row_factory = sqlite3.Row
        return dict(conn.execute("SELECT * FROM items WHERE paper_id=?", (paper_id,)).fetchone())


def test_unchanged_old_excel_preserves_new_database_notes(tmp_path):
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    library = LibraryService(tmp_path)
    library.update_personal_record(paper_id, {"user_notes": "库内独有关键词"})
    library.close()
    before = record(tmp_path, paper_id)
    assert service.refresh()["status"] == "success"
    after = record(tmp_path, paper_id)
    assert after["user_notes"] == "库内独有关键词"
    assert after["personal_updated_at"] == before["personal_updated_at"]


def test_sorted_rows_keep_six_personal_fields_and_can_be_searched(tmp_path):
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    edit(service, {"阅读进度": "待复读", "与课题的关系": "凝血悖论关联标记", "下一步": "核对图 2",
                   "个人思考": "一个判断", "个人理解程度": "还需复习", "用户笔记": "排序后仍属于第一篇"}, reverse=True)
    assert service.refresh()["status"] == "success"
    row = record(tmp_path, paper_id)
    assert row["reading_state"] == "待复读"
    assert row["user_notes"] == "排序后仍属于第一篇"
    assert row["personal_updated_at"]
    library = LibraryService(tmp_path)
    matches = library.search("凝血悖论关联标记")
    library.close()
    assert [m["paper_id"] for m in matches] == [paper_id]
    assert matches[0]["search_matches"][0]["content_type"] == "personal"


def test_conflicting_fields_preserve_workbook_and_both_values(tmp_path):
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    edit(service, {"用户笔记": "Excel 的新笔记"})
    library = LibraryService(tmp_path)
    library.update_personal_record(paper_id, {"user_notes": "对话里的新笔记"})
    library.close()
    original = service.target.read_bytes()
    result = service.refresh()
    assert result["status"] == "pending"
    assert result["error"]["code"] == "xlsx_field_conflict"
    assert service.target.read_bytes() == original
    assert record(tmp_path, paper_id)["user_notes"] == "对话里的新笔记"
    assert result["details"][0]["excel"] == "Excel 的新笔记"
    assert result["details"][0]["database"] == "对话里的新笔记"


def test_explicit_clear_and_identical_concurrent_change(tmp_path):
    paper_id = seed(tmp_path)[0]
    library = LibraryService(tmp_path)
    library.update_personal_record(paper_id, {"user_notes": "旧笔记"})
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    edit(service, {"用户笔记": None})
    assert service.refresh()["status"] == "success"
    assert record(tmp_path, paper_id)["user_notes"] == ""
    edit(service, {"用户笔记": "双方一致"})
    library.update_personal_record(paper_id, {"user_notes": "双方一致"})
    stamp = record(tmp_path, paper_id)["personal_updated_at"]
    assert service.refresh()["status"] == "success"
    assert record(tmp_path, paper_id)["personal_updated_at"] == stamp
    library.close()


def test_personal_record_whitelist_and_expected_value_guard(tmp_path):
    paper_id = seed(tmp_path)[0]
    library = LibraryService(tmp_path)
    with pytest.raises(ValueError, match="personal_fields_invalid"):
        library.update_personal_record(paper_id, {"title": "不能写系统字段"})
    with pytest.raises(ValueError, match="reading_state_invalid"):
        library.update_personal_record(paper_id, {"reading_state": "full_read_ready"})
    library.update_personal_record(paper_id, {"user_notes": "已修改"})
    with pytest.raises(ValueError, match="personal_record_conflict"):
        library.update_personal_record(paper_id, {"user_notes": "旧对话覆盖"}, expected={"user_notes": ""})
    library.close()


def legacy_workbook(service, paper_id, note):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "文献"
    sheet.append(("文献名", "文献 ID", "个人思考", "个人理解程度", "用户笔记", "旧手填占位列"))
    sheet.append(("旧论文", paper_id, "旧思考", "尚不理解", note, "此数据必须归档保留"))
    identity = book.create_sheet("_身份")
    identity.append(("row", "paper_id"))
    identity.append((2, paper_id))
    service.target.parent.mkdir(parents=True, exist_ok=True)
    book.save(service.target)
    book.close()


def test_legacy_upgrade_keeps_notes_and_archives_unmapped_content(tmp_path):
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    legacy_workbook(service, paper_id, "旧表长笔记")
    original = service.target.read_bytes()
    assert service.refresh()["status"] == "success"
    row = record(tmp_path, paper_id)
    assert row["personal_thoughts"] == "旧思考"
    assert row["user_notes"] == "旧表长笔记"
    assert row["reading_state"] == "未读"
    assert any(path.read_bytes() == original for path in (tmp_path / "backups" / "xlsx").glob("*.xlsx"))


def test_legacy_difference_does_not_overwrite_existing_database_notes(tmp_path):
    paper_id = seed(tmp_path)[0]
    library = LibraryService(tmp_path)
    library.update_personal_record(paper_id, {"user_notes": "库内较新笔记"})
    library.close()
    service = XlsxSnapshotService(tmp_path)
    legacy_workbook(service, paper_id, "旧 Excel 内容")
    original = service.target.read_bytes()
    assert service.refresh()["status"] == "pending"
    assert service.target.read_bytes() == original
    assert record(tmp_path, paper_id)["user_notes"] == "库内较新笔记"


def test_old_copy_and_backup_restore_keep_baselines(tmp_path):
    import shutil
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    old = service.target.read_bytes()
    edit(service, {"用户笔记": "第一次更改"})
    service.refresh()
    # 未编辑的旧副本不能覆盖后来导入的笔记。
    service.target.write_bytes(old)
    assert service.refresh()["status"] == "success"
    assert record(tmp_path, paper_id)["user_notes"] == "第一次更改"
    destination = tmp_path.parent / (tmp_path.name + "-restored")
    shutil.copytree(tmp_path, destination)
    restored = XlsxSnapshotService(destination)
    edit(restored, {"用户笔记": "换目录恢复后修改"})
    assert restored.refresh()["status"] == "success"
    assert record(destination, paper_id)["user_notes"] == "换目录恢复后修改"


def test_sync_receipt_updates_live_and_saved_environment_status(tmp_path):
    import json
    from scientific_reading.environment_status import EnvironmentStatusService
    seed(tmp_path)
    environment = EnvironmentStatusService(tmp_path)
    environment._write(environment.snapshot())
    assert environment.snapshot()["library"]["xlsx_pending"] == 2
    result = XlsxSnapshotService(tmp_path).refresh()
    saved = json.loads(environment.path.read_text(encoding="utf-8"))["library"]
    assert saved["xlsx_pending"] == 0
    assert saved["xlsx_last_export"]["export_id"] == result["export_id"]
    assert saved["xlsx_last_export"]["rows"] == 2
    assert environment.snapshot()["library"] == saved


def test_published_guide_and_confirmed_records_keep_separate_provenance(tmp_path):
    import json
    from scripts.reading_asset_fixture import seed_reading_assets
    fixture = seed_reading_assets(tmp_path, review_contract_version="full-review-v3")
    service = XlsxSnapshotService(tmp_path)
    records = service._review_rows()
    manifest = json.loads((tmp_path / fixture["reader"]).with_name("reader-manifest.json").read_text(encoding="utf-8"))
    expected = [entry["text"] for values in manifest["review"]["guide"].values() for entry in values]
    guides = [r for r in records if r["依据类型"] == "AI 导读"]
    assert expected and [r["内容"] for r in guides] == expected
    assert all(r["确认状态"] == "待核对" and r["block_ids"] for r in guides)
    assert any(r["确认状态"] == "历史确认 · 证据未核对" for r in records)
    with sqlite3.connect(tmp_path / "library.sqlite") as conn:
        conn.execute("UPDATE attachments SET sha256=?", ("d" * 64,))
    stale = service._review_rows()
    assert not any(r["依据类型"] == "AI 导读" for r in stale)
    assert any(r["确认状态"].startswith("历史确认") for r in stale)


def test_export_does_not_replace_a_workbook_saved_during_generation(tmp_path, monkeypatch):
    seed(tmp_path)
    service = XlsxSnapshotService(tmp_path)
    service.refresh()
    original = service._write_temp
    def concurrent_save(*args):
        path = original(*args)
        edit(service, {"用户笔记": "生成期间用户刚保存"})
        return path
    monkeypatch.setattr(service, "_write_temp", concurrent_save)
    assert service.refresh()["error"]["code"] == "xlsx_changed_during_export"
    monkeypatch.setattr(service, "_write_temp", original)
    assert service.refresh()["status"] == "success"
    book = openpyxl.load_workbook(service.target)
    headers = {c.value: c.column for c in book["文献"][1]}
    assert book["文献"].cell(2, headers["用户笔记"]).value == "生成期间用户刚保存"
    book.close()


def test_database_change_during_export_stays_pending(tmp_path, monkeypatch):
    paper_id = seed(tmp_path)[0]
    service = XlsxSnapshotService(tmp_path)
    original = service._write_temp
    def concurrent_update(*args):
        path = original(*args)
        library = LibraryService(tmp_path)
        library.update_personal_record(paper_id, {"user_notes": "生成期间新笔记"})
        library.close()
        return path
    monkeypatch.setattr(service, "_write_temp", concurrent_update)
    assert service.refresh()["status"] == "success"
    assert record(tmp_path, paper_id)["xlsx_sync_state"] == "pending"
    monkeypatch.setattr(service, "_write_temp", original)
    assert service.refresh()["status"] == "success"
    assert record(tmp_path, paper_id)["user_notes"] == "生成期间新笔记"


def test_parent_status_reads_completed_excel_child_instead_of_queued_snapshot(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from scientific_reading.__main__ import _job_foreground
    from scientific_reading.background_models import BackgroundRequest
    from scientific_reading.background_store import BackgroundJobStore
    from scientific_reading.foreground import ForegroundTimer
    from scientific_reading.reading_pipeline import ReadingPipeline
    paper_id = seed(tmp_path)[0]
    store = BackgroundJobStore(tmp_path)
    parent = store.create_or_get(BackgroundRequest(paper_id, "full_read_pipeline", "a" * 64, {}))
    child = store.create_or_get(BackgroundRequest(paper_id, "xlsx_snapshot", "b" * 64, {}))
    store.transition(child.job_id, "running")
    store.transition(child.job_id, "completed", result={"status": "success", "rows": 2})
    state = SimpleNamespace(stage_timings={}, stage_outputs={"schedule_derived_updates": {"status": "queued", "xlsx_job_id": child.job_id}})
    monkeypatch.setattr(ReadingPipeline, "_load", lambda *args, **kwargs: state)
    result, _ = _job_foreground(store, parent.job_id, ForegroundTimer())
    assert result.to_dict()["detail"]["xlsx"]["status"] == "completed"
    assert result.to_dict()["detail"]["xlsx"]["result"]["rows"] == 2
