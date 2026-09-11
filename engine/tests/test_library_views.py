import hashlib
import json

from openpyxl import load_workbook
import pytest

from scientific_reading.library_service import LibraryService
from scientific_reading.library_views import LibraryViews
from scientific_reading.models import PaperMetadata
from scientific_reading.xlsx_snapshot import XlsxSnapshotService


def setup(root):
    library = LibraryService(root)
    paper = library.ingest(PaperMetadata(title="Custom research fields", authors=[], abstract_en="The study used single cell sequencing."))["paper_id"]
    views = LibraryViews(library)
    field = views.field_save({"label": "研究焦点", "description": "与本人课题的具体联系"})["field_id"]
    return library, views, paper, field


def test_ai_evidence_manual_protection_and_stale_revision(tmp_path):
    library, views, paper, field = setup(tmp_path)
    abstract = library.get_item(paper)["abstract_en"]
    evidence = [{"kind": "abstract", "language": "en", "sha256": hashlib.sha256(abstract.encode()).hexdigest(), "quote": "single cell sequencing"}]
    request = {"paper_id": paper, "field_id": field, "value": "单细胞方法", "expected_revision": 0, "evidence": evidence}
    with pytest.raises(ValueError, match="research_evidence_required"):
        views.value_set({**request, "evidence": []}, origin="ai")
    result = views.value_set(request, origin="ai")["values"][0]
    assert (result["origin"], result["evidence"], result["revision"]) == ("ai", evidence, 1)
    with pytest.raises(ValueError, match="research_value_conflict"):
        views.value_set(request, origin="ai")
    views.value_set({**request, "value": "我自己的研究判断", "expected_revision": 1}, origin="manual")
    with pytest.raises(ValueError, match="research_manual_value_protected"):
        views.value_set({**request, "expected_revision": 2}, origin="ai")
    with pytest.raises(ValueError, match="research_abstract_evidence_changed"):
        views.value_set({**request, "evidence": [{**evidence[0], "sha256": "0" * 64}]}, origin="ai")
    library.close()


def test_scheme_names_reordering_and_excel_writeback_use_stable_ids(tmp_path):
    library, views, paper, field = setup(tmp_path)
    library.update_personal_record(paper, {"user_notes": "旧笔记保留"})
    views.value_set({"paper_id": paper, "field_id": field, "value": "旧研究记录", "expected_revision": 0}, origin="manual")
    scheme = views.scheme_save({"name": "课题视图", "config": {"columns": [
        {"field_id": "user_notes", "label": "我的笔记", "width": 42},
        {"field_id": field, "label": "课题重点", "width": 48, "wrap": False},
        {"field_id": "paper_id", "label": "永久编号", "hidden": True},
    ]}})["schemes"][0]
    views.scheme_default({"scheme_id": scheme["scheme_id"]})
    snapshot = XlsxSnapshotService(tmp_path)
    assert snapshot.refresh()["status"] == "success"
    book = load_workbook(snapshot.target)
    sheet = book["文献"]
    assert [cell.value for cell in sheet[1]][:3] == ["我的笔记", "课题重点", "永久编号"]
    assert sheet.column_dimensions["B"].width == 48 and not sheet["B2"].alignment.wrap_text
    assert sheet.column_dimensions["C"].hidden is True
    assert book["研究字段"]["D2"].value == "手动"
    sheet["A2"] = "重排后的笔记"
    sheet["B2"] = "重排后的课题记录"
    # 模拟整列重排：标题和每篇记录一起移动，稳定 ID 基线保持不变。
    for row in list(sheet.iter_rows()):
        values = [cell.value for cell in row][::-1]
        for cell, value in zip(row, values):
            cell.value = value
    book.save(snapshot.target)
    book.close()
    assert snapshot.import_user_fields()["updated"] == 1
    assert library.get_item(paper)["user_notes"] == "重排后的笔记"
    value = views.values(paper)["values"][0]
    assert value["value"] == "重排后的课题记录" and value["origin"] == "manual"
    assert snapshot.refresh()["status"] == "success"
    library.close()


def test_column_identity_tampering_and_custom_conflicts_preserve_file(tmp_path):
    library, views, paper, field = setup(tmp_path)
    snapshot = XlsxSnapshotService(tmp_path)
    assert snapshot.refresh()["status"] == "success"
    book = load_workbook(snapshot.target)
    cols = json.loads(book["_同步"]["B3"].value)
    cols[0]["field_id"] = "user_notes"
    book["_同步"]["B3"] = json.dumps(cols)
    book.save(snapshot.target)
    book.close()
    original = snapshot.target.read_bytes()
    assert snapshot.refresh()["error"]["code"] == "xlsx_column_identity_changed"
    assert snapshot.target.read_bytes() == original
    # 恢复真实列合同，然后制造同一研究字段的并发编辑。
    book = load_workbook(snapshot.target)
    cols[0]["field_id"] = "title"
    book["_同步"]["B3"] = json.dumps(cols)
    sheet = book["文献"]
    column = next(cell.column for cell in sheet[1] if cell.value == "研究焦点")
    sheet.cell(2, column).value = "Excel 中的新想法"
    book.save(snapshot.target)
    book.close()
    views.value_set({"paper_id": paper, "field_id": field, "value": "chat 中的新想法", "expected_revision": 0}, origin="manual")
    original = snapshot.target.read_bytes()
    result = snapshot.refresh()
    assert result["error"]["code"] == "xlsx_field_conflict"
    assert snapshot.target.read_bytes() == original
    assert views.values(paper)["values"][0]["value"] == "chat 中的新想法"
    library.close()


def test_category_view_and_v2_notes_migration(tmp_path):
    library, views, paper, _field = setup(tmp_path)
    folder = library.create_folder("我的课题")["folder_id"]
    library.move_items([paper], folder)
    library.ingest(PaperMetadata(title="Other category", authors=[]))
    scheme = views.scheme_save({"scope_folder_id": folder, "name": "分类方案", "config": {"columns": [{"field_id": "title", "label": "课题文献"}]}})["schemes"][0]
    views.scheme_default({"scope_folder_id": folder, "scheme_id": scheme["scheme_id"]})
    category = XlsxSnapshotService(tmp_path, folder)
    assert category.refresh()["rows"] == 1
    book = load_workbook(category.target)
    assert book["文献"]["A1"].value == "课题文献"
    book.close()
    global_view = XlsxSnapshotService(tmp_path)
    assert global_view.refresh()["rows"] == 2
    book = load_workbook(global_view.target)
    assert book["文献"]["A1"].value == "文献名"
    book["_同步"]["B1"] = "xlsx-library-v2"
    export_id = book["_同步"]["B2"].value
    baseline = json.loads(library.conn.execute("SELECT baseline_json FROM xlsx_exports WHERE export_id=?", (export_id,)).fetchone()[0])
    library.conn.execute("UPDATE xlsx_exports SET baseline_json=? WHERE export_id=?", (json.dumps(baseline["rows"]), export_id))
    library.conn.commit()
    col = next(cell.column for cell in book["文献"][1] if cell.value == "用户笔记")
    book["文献"].cell(2, col).value = "v0.1 旧工作簿里的笔记"
    book.save(global_view.target)
    book.close()
    assert global_view.refresh()["status"] == "success"
    assert library.get_item(paper)["user_notes"] == "v0.1 旧工作簿里的笔记"
    assert len(list((tmp_path / "backups" / "xlsx").glob("*.xlsx"))) >= 1
    library.close()
