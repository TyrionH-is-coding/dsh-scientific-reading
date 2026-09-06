import openpyxl
import pytest

from scientific_reading.background_models import BackgroundRequest
from scientific_reading.background_store import BackgroundJobStore
from scientific_reading.library_service import LibraryService
from scientific_reading.models import PaperMetadata
from scientific_reading.worker import run_job
from scientific_reading.xlsx_snapshot import XlsxSnapshotService, XLSX_COLUMNS


def test_xlsx_worker_waits_for_identity_repair_then_resumes(tmp_path):
    library = LibraryService(tmp_path)
    paper_id = library.ingest(PaperMetadata(title="Offline bridge paper", doi="10.5555/xlsx-worker"))["paper_id"]
    library.close()
    snapshot = XlsxSnapshotService(tmp_path)
    snapshot.refresh()
    workbook = openpyxl.load_workbook(snapshot.target)
    identity_column = XLSX_COLUMNS.index("文献 ID") + 1
    notes_column = XLSX_COLUMNS.index("用户笔记") + 1
    workbook["文献"].cell(2, identity_column, "changed_identity")
    workbook["文献"].cell(2, notes_column, "等待修复的笔记")
    workbook.save(snapshot.target)
    workbook.close()
    original = snapshot.target.read_bytes()
    store = BackgroundJobStore(tmp_path)
    request = BackgroundRequest(paper_id, "xlsx_snapshot", "a" * 64, {"data_root": str(tmp_path)})
    job_id = store.create_or_get(request).job_id

    assert run_job(store, job_id) == 2

    status = store.load_status(job_id)
    assert status.state == "waiting_user"
    assert status.reason_code == "xlsx_identity_conflict"
    assert status.required_input["conflicts"] == 1
    assert snapshot.target.read_bytes() == original
    workbook = openpyxl.load_workbook(snapshot.target)
    workbook["文献"].cell(2, identity_column, paper_id)
    workbook.save(snapshot.target)
    workbook.close()

    store.transition(job_id, "queued")
    assert run_job(store, job_id) == 0
    assert store.load_status(job_id).state == "completed"
    workbook = openpyxl.load_workbook(snapshot.target)
    assert workbook["文献"].cell(2, notes_column).value == "等待修复的笔记"
    workbook.close()


@pytest.mark.parametrize("error, exit_code, state, code", [
    (PermissionError("locked"), 2, "waiting_user", "xlsx_replace_permission_denied"),
    (OSError("write failed"), 4, "failed", "xlsx_snapshot_failed"),
])
def test_xlsx_worker_reports_write_failure_and_keeps_original(tmp_path, monkeypatch, error, exit_code, state, code):
    library = LibraryService(tmp_path)
    paper_id = library.ingest(PaperMetadata(title="Offline write failure", doi="10.5555/xlsx-write"))["paper_id"]
    library.close()
    snapshot = XlsxSnapshotService(tmp_path)
    snapshot.refresh()
    original = snapshot.target.read_bytes()
    store = BackgroundJobStore(tmp_path)
    request = BackgroundRequest(paper_id, "xlsx_snapshot", "b" * 64, {"data_root": str(tmp_path)})
    job_id = store.create_or_get(request).job_id
    real_replace = __import__("os").replace

    def fail_snapshot_replace(source, target):
        if str(target) == str(snapshot.target):
            raise error
        return real_replace(source, target)

    monkeypatch.setattr("scientific_reading.xlsx_snapshot.os.replace", fail_snapshot_replace)

    assert run_job(store, job_id) == exit_code
    status = store.load_status(job_id)
    assert status.state == state
    assert (status.reason_code if state == "waiting_user" else status.error) == code
    assert snapshot.target.read_bytes() == original
